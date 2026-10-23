"""
pptx_creator.py — JARVIS plugin

Creates a designed PowerPoint (.pptx) from slides Gemini extracts from the
user's request. Unlike a one-template-fits-all deck, every slide picks a
LAYOUT suited to its content (title, section divider, bullets, two-column,
table, native chart, stat callout, quote, process flow, timeline, full-bleed
image, closing) and a BACKGROUND variant (primary / muted / accent) — so a
10-slide deck actually looks like 10 considered slides, not one slide copied
ten times, while every slide still shares one consistent colour/font theme.

Real PowerPoint tooling used, not just textboxes:
  - native pptx tables (add_table)
  - native pptx charts — bar/column/line/pie (add_chart)
  - shape-built process & timeline diagrams (ovals, connectors)
  - image cropping to fill a frame without distortion (Picture.crop_*)
  - semi-transparent overlay panels (alpha on solid fills)
  - speaker notes, slide numbering / footer, section "kicker" labels

Requires (all already in requirements.txt): python-pptx, ddgs, requests, pillow.

Drop this file into /plugins/ — no other file needs to change.
"""

import io
from pathlib import Path
from datetime import datetime

PLUGIN = {
    "name": "create_powerpoint",
    "description": (
        "Creates a new, professionally designed PowerPoint (.pptx) and saves it "
        "to disk. Use when the user asks to 'make a powerpoint', 'create a "
        "presentation', 'build slides', 'put this into a pptx', or similar. "
        "This tool is layout-aware: give EACH slide the 'layout' that actually "
        "fits its content (a stat gets 'stat_highlight', a comparison gets "
        "'two_column' or 'table', data gets a real native 'chart', a topic "
        "change gets a 'section' divider) rather than putting everything on a "
        "generic bullets slide — a good deck mixes several layouts. All layouts "
        "share one colour/font theme so the deck still looks cohesive. "
        "Do NOT use this to read, summarize, or extract text from an EXISTING "
        "pptx file — that's handled by the file processor instead."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "title": {
                "type": "STRING",
                "description": "Title slide headline; also the default file name.",
            },
            "subtitle": {
                "type": "STRING",
                "description": "Optional subtitle under the title.",
            },
            "author": {
                "type": "STRING",
                "description": "Optional presenter/author name shown on the title slide.",
            },
            "template": {
                "type": "STRING",
                "enum": ["modern_dark", "minimal_light", "corporate_blue",
                         "elegant_gold", "bold_contrast"],
                "description": (
                    "Colour/font theme applied to every slide (the LAYOUT varies "
                    "per slide, the theme stays constant). 'modern_dark' (default) — "
                    "dark bg, electric-cyan + magenta accents, sleek/futuristic. "
                    "'minimal_light' — white bg, grey + blue accents, clean. "
                    "'corporate_blue' — navy, blue + amber accents, business-safe. "
                    "'elegant_gold' — charcoal, gold + violet accents, upscale. "
                    "'bold_contrast' — near-black, orange + teal accents, high energy."
                ),
            },
            "slides": {
                "type": "ARRAY",
                "description": (
                    "The slides after the title slide, in order. Pick a 'layout' per "
                    "slide that fits its content — mixing layouts is the point."
                ),
                "items": {
                    "type": "OBJECT",
                    "properties": {
                        "layout": {
                            "type": "STRING",
                            "enum": ["bullets", "two_column", "table", "chart",
                                     "stat_highlight", "quote", "process", "timeline",
                                     "image_full", "section"],
                            "description": (
                                "'bullets' — heading + bullet list, optional image. "
                                "'two_column' — two headed bullet lists side by side "
                                "(comparisons, pros/cons). "
                                "'table' — a real PowerPoint table from 'table'. "
                                "'chart' — a real native PowerPoint chart from 'chart' "
                                "(bar/column/line/pie) — use this for numbers, not bullets. "
                                "'stat_highlight' — 1-3 big numbers with labels, for the "
                                "single most important metric(s) on a slide. "
                                "'quote' — one big quotation, centered, for a testimonial "
                                "or key statement. "
                                "'process' — up to 5 numbered steps in a flow. "
                                "'timeline' — up to 6 dated/ordered events on a line. "
                                "'image_full' — full-bleed photo with heading text overlaid. "
                                "'section' — big divider slide marking a new topic/section — "
                                "use this between subject changes instead of jumping "
                                "straight from one bullets slide to an unrelated one."
                            ),
                        },
                        "background": {
                            "type": "STRING",
                            "enum": ["primary", "muted", "accent"],
                            "description": (
                                "Optional override for this slide's background variant "
                                "within the theme — 'primary' (theme's normal bg), 'muted' "
                                "(a subtly shifted shade, good for alternating slides so "
                                "back-to-back slides don't look identical), or 'accent' "
                                "(full theme accent colour, high impact — used "
                                "automatically for section/quote/closing if not set). "
                                "Leave unset to use the layout's sensible default."
                            ),
                        },
                        "kicker": {
                            "type": "STRING",
                            "description": (
                                "Optional small eyebrow label above the heading, e.g. "
                                "'MARKET OVERVIEW' or 'STEP 2 OF 4' — helps signal which "
                                "topic/section this slide belongs to."
                            ),
                        },
                        "heading": {"type": "STRING", "description": "Slide heading/title."},
                        "subheading": {
                            "type": "STRING",
                            "description": "Optional supporting line under the heading "
                                            "(used by 'section' and 'image_full').",
                        },
                        "bullets": {
                            "type": "ARRAY",
                            "description": "Bullet points, plain strings. Used by 'bullets'.",
                            "items": {"type": "STRING"},
                        },
                        "left_heading": {"type": "STRING", "description": "'two_column' left header."},
                        "left_bullets": {"type": "ARRAY", "items": {"type": "STRING"},
                                          "description": "'two_column' left bullet list."},
                        "right_heading": {"type": "STRING", "description": "'two_column' right header."},
                        "right_bullets": {"type": "ARRAY", "items": {"type": "STRING"},
                                           "description": "'two_column' right bullet list."},
                        "table": {
                            "type": "OBJECT",
                            "description": "For layout 'table'.",
                            "properties": {
                                "headers": {"type": "ARRAY", "items": {"type": "STRING"}},
                                "rows": {
                                    "type": "ARRAY",
                                    "description": "Array of row-arrays of strings.",
                                    "items": {"type": "ARRAY", "items": {"type": "STRING"}},
                                },
                            },
                        },
                        "chart": {
                            "type": "OBJECT",
                            "description": "For layout 'chart' — a real native PowerPoint chart.",
                            "properties": {
                                "type": {"type": "STRING", "enum": ["bar", "column", "line", "pie"]},
                                "categories": {"type": "ARRAY", "items": {"type": "STRING"}},
                                "series": {
                                    "type": "ARRAY",
                                    "description": "One entry per data series (pie charts: exactly one).",
                                    "items": {
                                        "type": "OBJECT",
                                        "properties": {
                                            "name": {"type": "STRING"},
                                            "values": {"type": "ARRAY", "items": {"type": "NUMBER"}},
                                        },
                                        "required": ["values"],
                                    },
                                },
                            },
                        },
                        "stats": {
                            "type": "ARRAY",
                            "description": "For 'stat_highlight' — 1 to 3 entries.",
                            "items": {
                                "type": "OBJECT",
                                "properties": {
                                    "value": {"type": "STRING", "description": "e.g. '87%', '$4.2M', '12x'"},
                                    "label": {"type": "STRING"},
                                },
                                "required": ["value", "label"],
                            },
                        },
                        "quote_text": {"type": "STRING", "description": "For 'quote'."},
                        "quote_author": {"type": "STRING", "description": "For 'quote', optional attribution."},
                        "process_steps": {
                            "type": "ARRAY",
                            "description": "For 'process' — up to 5, in order.",
                            "items": {
                                "type": "OBJECT",
                                "properties": {
                                    "title": {"type": "STRING"},
                                    "description": {"type": "STRING"},
                                },
                                "required": ["title"],
                            },
                        },
                        "timeline_events": {
                            "type": "ARRAY",
                            "description": "For 'timeline' — up to 6, in order.",
                            "items": {
                                "type": "OBJECT",
                                "properties": {
                                    "date": {"type": "STRING"},
                                    "label": {"type": "STRING"},
                                },
                                "required": ["label"],
                            },
                        },
                        "image_query": {
                            "type": "STRING",
                            "description": "Short search term; JARVIS fetches a usable web image. "
                                            "Used by 'bullets' (optional side image) and 'image_full'.",
                        },
                        "image_path": {
                            "type": "STRING",
                            "description": "Local image path already on disk, used instead of searching.",
                        },
                        "notes": {
                            "type": "STRING",
                            "description": "Optional speaker notes attached to this slide.",
                        },
                    },
                    "required": ["layout", "heading"],
                },
            },
            "closing": {
                "type": "OBJECT",
                "description": "Optional closing/thank-you slide, added last.",
                "properties": {
                    "heading": {"type": "STRING"},
                    "subheading": {"type": "STRING"},
                },
            },
            "output_path": {
                "type": "STRING",
                "description": (
                    "Optional. Full file path or bare filename for the saved .pptx. "
                    "If omitted, JARVIS saves to the Desktop using a slugified version "
                    "of the title with a timestamp."
                ),
            },
        },
        "required": ["title", "slides"],
    },
}

_MAX_IMAGE_BYTES = 8 * 1024 * 1024   # 8 MB safety cap per image
_IMG_TIMEOUT = 10                    # seconds
_CACHE_DIR = Path(__file__).resolve().parent.parent / "cache" / "pptx_images"

# Each theme: bg (base background), accent (primary highlight), accent2
# (secondary highlight — multi-series charts, alternating emphasis),
# title/body text colours, font family.
THEMES = {
    "modern_dark": {
        "bg": (0x0B, 0x0F, 0x19), "accent": (0x00, 0xE5, 0xFF), "accent2": (0xFF, 0x2E, 0x92),
        "title": (0xFF, 0xFF, 0xFF), "body": (0xC9, 0xD3, 0xDF), "font": "Segoe UI",
    },
    "minimal_light": {
        "bg": (0xFA, 0xFA, 0xFA), "accent": (0x2F, 0x80, 0xED), "accent2": (0x8A, 0x8F, 0x98),
        "title": (0x1A, 0x1A, 0x1A), "body": (0x40, 0x40, 0x40), "font": "Calibri",
    },
    "corporate_blue": {
        "bg": (0x0E, 0x22, 0x3E), "accent": (0x3D, 0x8B, 0xFF), "accent2": (0xFF, 0xB4, 0x3D),
        "title": (0xFF, 0xFF, 0xFF), "body": (0xD5, 0xE2, 0xF0), "font": "Calibri",
    },
    "elegant_gold": {
        "bg": (0x1C, 0x1B, 0x19), "accent": (0xD4, 0xAF, 0x37), "accent2": (0x8A, 0x5C, 0xF6),
        "title": (0xF5, 0xF0, 0xE6), "body": (0xCE, 0xC6, 0xB8), "font": "Georgia",
    },
    "bold_contrast": {
        "bg": (0x12, 0x10, 0x10), "accent": (0xFF, 0x6A, 0x1A), "accent2": (0x00, 0xC2, 0xA8),
        "title": (0xFF, 0xFF, 0xFF), "body": (0xE8, 0xE2, 0xDD), "font": "Segoe UI",
    },
}


# ── small generic helpers ───────────────────────────────────────────────────

def _slugify(text: str) -> str:
    keep = [c if c.isalnum() else "_" for c in text.strip().lower()]
    slug = "".join(keep)
    while "__" in slug:
        slug = slug.replace("__", "_")
    return slug.strip("_") or "presentation"


def _resolve_output_path(output_path: str, title: str) -> Path:
    if output_path:
        p = Path(output_path).expanduser()
        if p.suffix.lower() != ".pptx":
            p = p.with_suffix(".pptx")
        p.parent.mkdir(parents=True, exist_ok=True)
        return p
    desktop = Path.home() / "Desktop"
    desktop.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return desktop / f"{_slugify(title)}_{stamp}.pptx"


def _luminance(rgb) -> float:
    r, g, b = [c / 255 for c in rgb]
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _is_light(rgb) -> bool:
    return _luminance(rgb) > 0.6


def _contrast_text(rgb):
    """Black or white — whichever reads better on top of rgb."""
    return (0x1A, 0x1A, 0x1A) if _is_light(rgb) else (0xFF, 0xFF, 0xFF)


def _shade(rgb, amount: float):
    """amount > 0 lightens toward white, amount < 0 darkens toward black. Roughly -100..100."""
    r, g, b = rgb
    if amount >= 0:
        r = r + (255 - r) * amount / 100
        g = g + (255 - g) * amount / 100
        b = b + (255 - b) * amount / 100
    else:
        f = max(0.0, 1 + amount / 100)
        r, g, b = r * f, g * f, b * f
    return (int(max(0, min(255, r))), int(max(0, min(255, g))), int(max(0, min(255, b))))


def _bg_variant(theme: dict, variant: str):
    """primary / muted / accent — the background-picking mechanism behind the
    per-slide 'background' field."""
    if variant == "accent":
        return theme["accent"]
    if variant == "muted":
        # push the base background slightly toward its opposite pole so it
        # visibly differs from a plain 'primary' slide without breaking theme
        return _shade(theme["bg"], 8 if not _is_light(theme["bg"]) else -6)
    return theme["bg"]


# ── image fetching (unchanged behaviour) ────────────────────────────────────

def _fetch_image_url(query: str):
    """Search the web for one usable image URL. Returns None on any failure."""
    try:
        try:
            from ddgs import DDGS
        except ImportError:
            from duckduckgo_search import DDGS
        with DDGS() as ddgs:
            for r in ddgs.images(query, max_results=8):
                url = r.get("image") or r.get("url") or ""
                if url.lower().split("?")[0].endswith((".jpg", ".jpeg", ".png", ".webp")):
                    return url
        return None
    except Exception:
        return None


def _download_image(url: str):
    """Download + validate an image URL, normalize to JPEG, cache it. None on failure."""
    try:
        import requests
        from PIL import Image

        _CACHE_DIR.mkdir(parents=True, exist_ok=True)
        resp = requests.get(url, timeout=_IMG_TIMEOUT, stream=True,
                             headers={"User-Agent": "Mozilla/5.0"})
        resp.raise_for_status()
        ctype = resp.headers.get("Content-Type", "")
        if "svg" in ctype:
            return None

        buf = io.BytesIO()
        size = 0
        for chunk in resp.iter_content(8192):
            size += len(chunk)
            if size > _MAX_IMAGE_BYTES:
                return None
            buf.write(chunk)
        buf.seek(0)

        img = Image.open(buf)
        img.load()
        img = img.convert("RGB")

        dest = _CACHE_DIR / f"img_{abs(hash(url)) % 10**10}.jpg"
        img.save(dest, "JPEG", quality=88)
        return dest
    except Exception:
        return None


def _resolve_slide_image(entry: dict):
    """image_path (local, wins if present) or image_query (web search) -> local Path."""
    image_path = (entry.get("image_path") or "").strip()
    if image_path:
        p = Path(image_path).expanduser()
        return p if p.is_file() else None

    image_query = (entry.get("image_query") or "").strip()
    if image_query:
        url = _fetch_image_url(image_query)
        if url:
            return _download_image(url)
    return None


# ── pptx primitives (all pptx imports stay local — plugin must import cleanly
#    even if python-pptx isn't installed yet; run() checks for it up front) ──

def _rgb(t):
    from pptx.dml.color import RGBColor
    return RGBColor(*t)


def _set_fill_alpha(shape, alpha_pct: int):
    """Make a solid-fill shape semi-transparent. alpha_pct: 0 (invisible) - 100 (opaque).
    python-pptx has no first-class API for this; it's a documented XML add-on."""
    try:
        from pptx.oxml.ns import qn
        sp = shape.fill.fore_color._xFill
        srgb = sp.find(qn("a:srgbClr"))
        if srgb is None:
            return
        for existing in srgb.findall(qn("a:alpha")):
            srgb.remove(existing)
        alpha_el = sp.makeelement(qn("a:alpha"), {"val": str(int(alpha_pct * 1000))})
        srgb.append(alpha_el)
    except Exception:
        pass  # cosmetic only — never let a transparency tweak break the deck


def _set_background(slide, rgb):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(rgb)


def _add_shape(slide, shape_enum, left, top, width, height, fill_rgb=None, line_rgb=None, line_pt=None):
    shape = slide.shapes.add_shape(shape_enum, left, top, width, height)
    if fill_rgb is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill_rgb)
    else:
        shape.fill.background()
    if line_rgb is not None:
        from pptx.util import Pt
        shape.line.color.rgb = _rgb(line_rgb)
        shape.line.width = Pt(line_pt or 1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _add_accent_bar(slide, prs, rgb, vertical=True):
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE
    if vertical:
        bar = _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height, fill_rgb=rgb)
    else:
        bar = _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12), fill_rgb=rgb)
    return bar


def _add_textbox(slide, left, top, width, height, text, size, color, font,
                  bold=False, align="left", italic=False):
    from pptx.util import Pt
    from pptx.enum.text import PP_ALIGN
    align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font
    p.font.color.rgb = _rgb(color)
    p.alignment = align_map.get(align, PP_ALIGN.LEFT)
    return box


def _add_kicker(slide, theme, text, left, top, width, align="left"):
    if not text:
        return
    _add_textbox(slide, left, top, width, _in(0.4), text.upper(), 12,
                 theme["accent"], theme["font"], bold=True, align=align)


def _add_footer(slide, prs, theme, deck_title, idx, total):
    from pptx.util import Pt
    footer_color = _shade(theme["body"], -10 if _is_light(theme["bg"]) else -20)
    _add_textbox(slide, _in(0.5), prs.slide_height - _in(0.45), _in(6), _in(0.35),
                 deck_title, 10, footer_color, theme["font"])
    _add_textbox(slide, prs.slide_width - _in(1.5), prs.slide_height - _in(0.45),
                 _in(1.0), _in(0.35), f"{idx:02d} / {total:02d}", 10,
                 footer_color, theme["font"], align="right")


def _in(v: float):
    from pptx.util import Inches
    return Inches(v)


def _pt(v: float):
    from pptx.util import Pt
    return Pt(v)


def _i(v) -> int:
    """Round an EMU-arithmetic result (may have become a plain float) back to
    a whole-EMU int, which is what pptx's XML writers require."""
    return int(round(v))


def _blank_slide(prs, theme, bg_rgb):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    _set_background(slide, bg_rgb)
    return slide


# ── per-layout slide builders ────────────────────────────────────────────────

def _build_title_slide(prs, theme, title, subtitle, author):
    slide = _blank_slide(prs, theme, theme["bg"])
    _add_accent_bar(slide, prs, theme["accent"], vertical=False)
    from pptx.enum.shapes import MSO_SHAPE
    _add_shape(slide, MSO_SHAPE.OVAL, prs.slide_width - _in(2.6), -_in(1.2), _in(3.2), _in(3.2),
               fill_rgb=_shade(theme["bg"], 6 if not _is_light(theme["bg"]) else -4))

    _add_textbox(slide, _in(0.8), _in(2.5), prs.slide_width - _in(1.6), _in(1.7),
                 title, 40, theme["title"], theme["font"], bold=True)
    if subtitle:
        _add_textbox(slide, _in(0.8), _in(3.9), prs.slide_width - _in(1.6), _in(0.9),
                     subtitle, 18, theme["accent"], theme["font"])
    if author:
        _add_textbox(slide, _in(0.8), prs.slide_height - _in(0.9), _in(6), _in(0.5),
                     author, 13, theme["body"], theme["font"])
    return slide


def _build_section_slide(prs, theme, heading, subheading, background, kicker):
    from pptx.enum.shapes import MSO_SHAPE
    bg = _bg_variant(theme, background or "accent")
    slide = _blank_slide(prs, theme, bg)
    fg_title = _contrast_text(bg) if background == "accent" or (not background) else theme["title"]
    fg_body = fg_title if (background == "accent" or not background) else theme["body"]

    _add_shape(slide, MSO_SHAPE.RECTANGLE,
               0, prs.slide_height - _in(0.16), prs.slide_width, _in(0.16),
               fill_rgb=_shade(bg, -15) if _is_light(bg) else _shade(bg, 25))

    if kicker:
        _add_kicker(slide, {**theme, "accent": fg_title}, kicker, _in(0.9), _in(2.1), _in(8))
    _add_textbox(slide, _in(0.9), _in(2.6), prs.slide_width - _in(1.8), _in(1.8),
                 heading, 44, fg_title, theme["font"], bold=True)
    if subheading:
        _add_textbox(slide, _in(0.9), _in(3.9), prs.slide_width - _in(1.8), _in(1.0),
                     subheading, 18, fg_body, theme["font"])
    return slide


def _build_bullets_slide(prs, theme, heading, bullets, image_path, background, kicker):
    bg = _bg_variant(theme, background or "primary")
    slide = _blank_slide(prs, theme, bg)
    _add_accent_bar(slide, prs, theme["accent"], vertical=True)

    has_image = image_path is not None
    text_width = (prs.slide_width - _in(1.0) - _in(4.6)) if has_image else (prs.slide_width - _in(1.0))
    top = _in(0.5)
    if kicker:
        _add_kicker(slide, theme, kicker, _in(0.6), top, text_width)
        top = top + _in(0.45)

    _add_textbox(slide, _in(0.6), top, text_width, _in(0.9),
                 heading, 28, theme["title"], theme["font"], bold=True)

    body_top = top + _in(1.0)
    body_box = slide.shapes.add_textbox(_in(0.6), body_top, text_width,
                                         prs.slide_height - body_top - _in(0.6))
    tf = body_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"›  {bullet}"
        p.font.size = _pt(18)
        p.font.name = theme["font"]
        p.font.color.rgb = _rgb(theme["body"])
        p.space_after = _pt(10)

    if has_image:
        _place_cropped_image(slide, image_path,
                              prs.slide_width - _in(4.6), body_top,
                              _in(4.0), prs.slide_height - body_top - _in(0.6))
    return slide


def _build_two_column_slide(prs, theme, heading, l_head, l_bullets, r_head, r_bullets, background, kicker):
    from pptx.enum.shapes import MSO_SHAPE
    bg = _bg_variant(theme, background or "primary")
    slide = _blank_slide(prs, theme, bg)
    _add_accent_bar(slide, prs, theme["accent"], vertical=False)

    top = _in(0.5)
    if kicker:
        _add_kicker(slide, theme, kicker, _in(0.6), top, prs.slide_width - _in(1.2))
        top = top + _in(0.45)
    _add_textbox(slide, _in(0.6), top, prs.slide_width - _in(1.2), _in(0.8),
                 heading, 28, theme["title"], theme["font"], bold=True)

    col_top = top + _in(1.0)
    col_w = int((prs.slide_width - _in(1.6)) / 2)
    divider_x = int(_in(0.6) + col_w + _in(0.2))
    _add_shape(slide, MSO_SHAPE.RECTANGLE,
               divider_x, col_top, _pt(1.5), prs.slide_height - col_top - _in(0.5),
               fill_rgb=_shade(theme["body"], -30 if not _is_light(bg) else 30))

    for col_idx, (head, bullets, x) in enumerate([
        (l_head, l_bullets, _in(0.6)),
        (r_head, r_bullets, divider_x + _in(0.4)),
    ]):
        accent = theme["accent"] if col_idx == 0 else theme["accent2"]
        if head:
            _add_textbox(slide, x, col_top, col_w - _in(0.3), _in(0.5),
                         head, 18, accent, theme["font"], bold=True)
        box = slide.shapes.add_textbox(x, col_top + _in(0.55), col_w - _in(0.3),
                                        prs.slide_height - col_top - _in(1.1))
        tf = box.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(bullets or []):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"›  {bullet}"
            p.font.size = _pt(16)
            p.font.name = theme["font"]
            p.font.color.rgb = _rgb(theme["body"])
            p.space_after = _pt(8)
    return slide


def _build_table_slide(prs, theme, heading, table_spec, background, kicker):
    bg = _bg_variant(theme, background or "primary")
    slide = _blank_slide(prs, theme, bg)
    _add_accent_bar(slide, prs, theme["accent"], vertical=True)

    top = _in(0.5)
    if kicker:
        _add_kicker(slide, theme, kicker, _in(0.6), top, prs.slide_width - _in(1.2))
        top = top + _in(0.45)
    _add_textbox(slide, _in(0.6), top, prs.slide_width - _in(1.2), _in(0.8),
                 heading, 28, theme["title"], theme["font"], bold=True)

    headers = (table_spec or {}).get("headers") or []
    rows = (table_spec or {}).get("rows") or []
    n_cols = max(1, len(headers) or (len(rows[0]) if rows else 1))
    n_rows = len(rows) + (1 if headers else 0)
    if n_rows == 0:
        return slide

    table_top = top + _in(1.0)
    gframe = slide.shapes.add_table(n_rows, n_cols, _in(0.6), table_top,
                                     prs.slide_width - _in(1.2),
                                     prs.slide_height - table_top - _in(0.6))
    tbl = gframe.table

    header_text_rgb = _contrast_text(theme["accent"])
    row_offset = 0
    if headers:
        for c in range(n_cols):
            cell = tbl.cell(0, c)
            cell.text = str(headers[c]) if c < len(headers) else ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(theme["accent"])
            for p in cell.text_frame.paragraphs:
                p.font.bold = True
                p.font.size = _pt(14)
                p.font.name = theme["font"]
                p.font.color.rgb = _rgb(header_text_rgb)
        row_offset = 1

    for r, row in enumerate(rows):
        zebra = _shade(bg, 8 if not _is_light(bg) else -6) if r % 2 else _shade(bg, 3 if not _is_light(bg) else -2)
        for c in range(n_cols):
            cell = tbl.cell(r + row_offset, c)
            cell.text = str(row[c]) if c < len(row) else ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(zebra)
            for p in cell.text_frame.paragraphs:
                p.font.size = _pt(13)
                p.font.name = theme["font"]
                p.font.color.rgb = _rgb(theme["body"])
    return slide


def _build_chart_slide(prs, theme, heading, chart_spec, background, kicker):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    bg = _bg_variant(theme, background or "primary")
    slide = _blank_slide(prs, theme, bg)
    _add_accent_bar(slide, prs, theme["accent"], vertical=True)

    top = _in(0.5)
    if kicker:
        _add_kicker(slide, theme, kicker, _in(0.6), top, prs.slide_width - _in(1.2))
        top = top + _in(0.45)
    _add_textbox(slide, _in(0.6), top, prs.slide_width - _in(1.2), _in(0.8),
                 heading, 28, theme["title"], theme["font"], bold=True)

    chart_spec = chart_spec or {}
    categories = chart_spec.get("categories") or []
    series_list = chart_spec.get("series") or []
    if not categories or not series_list:
        _add_textbox(slide, _in(0.6), top + _in(1.2), prs.slide_width - _in(1.2), _in(1.0),
                     "(no chart data supplied)", 16, theme["body"], theme["font"])
        return slide

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for s in series_list:
        chart_data.add_series(s.get("name") or "Series", s.get("values") or [])

    kind = (chart_spec.get("type") or "column").lower()
    xl_type_map = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }
    xl_type = xl_type_map.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)

    chart_top = top + _in(1.0)
    gframe = slide.shapes.add_chart(
        xl_type, _in(0.7), chart_top,
        prs.slide_width - _in(1.4), prs.slide_height - chart_top - _in(0.6),
        chart_data,
    )
    chart = gframe.chart
    palette = [theme["accent"], theme["accent2"], _shade(theme["accent"], -20), _shade(theme["accent2"], -20)]

    try:
        if xl_type == XL_CHART_TYPE.PIE:
            chart.has_legend = True
            chart.legend.include_in_layout = False
            points = chart.plots[0].series[0].points
            for i, pt in enumerate(points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
        else:
            chart.has_legend = len(series_list) > 1
            if chart.has_legend:
                chart.legend.include_in_layout = False
                chart.legend.font.size = _pt(11)
                chart.legend.font.color.rgb = _rgb(theme["body"])
            for i, series in enumerate(chart.plots[0].series):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
                if xl_type == XL_CHART_TYPE.LINE_MARKERS:
                    series.format.line.color.rgb = _rgb(palette[i % len(palette)])
            for axis in (chart.category_axis, chart.value_axis):
                axis.tick_labels.font.size = _pt(11)
                axis.tick_labels.font.color.rgb = _rgb(theme["body"])
                axis.format.line.color.rgb = _rgb(_shade(theme["body"], -20))
    except Exception:
        pass  # theme styling is cosmetic — a plain default chart still beats no chart
    return slide


def _build_stat_slide(prs, theme, heading, stats, background, kicker):
    bg = _bg_variant(theme, background or "muted")
    slide = _blank_slide(prs, theme, bg)
    _add_accent_bar(slide, prs, theme["accent"], vertical=False)

    top = _in(0.5)
    if kicker:
        _add_kicker(slide, theme, kicker, _in(0.6), top, prs.slide_width - _in(1.2))
        top = top + _in(0.45)
    if heading:
        _add_textbox(slide, _in(0.6), top, prs.slide_width - _in(1.2), _in(0.7),
                     heading, 24, theme["title"], theme["font"], bold=True)

    stats = (stats or [])[:3]
    if not stats:
        return slide
    n = len(stats)
    card_w = _i((prs.slide_width - _in(1.2) - _in(0.4) * (n - 1)) / n)
    card_top = _in(2.6)
    card_h = _in(2.6)
    palette = [theme["accent"], theme["accent2"], _shade(theme["accent"], -15)]

    for i, stat in enumerate(stats):
        x = _i(_in(0.6) + i * (card_w + _in(0.4)))
        from pptx.enum.shapes import MSO_SHAPE
        card = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, card_top, card_w, card_h,
                           fill_rgb=_shade(bg, 10 if not _is_light(bg) else -6))
        try:
            card.adjustments[0] = 0.06
        except Exception:
            pass
        _add_textbox(slide, x + _in(0.15), card_top + _in(0.35), card_w - _in(0.3), _in(1.1),
                     str(stat.get("value", "")), 40, palette[i % len(palette)], theme["font"],
                     bold=True, align="center")
        _add_textbox(slide, x + _in(0.15), card_top + _in(1.6), card_w - _in(0.3), _in(0.8),
                     str(stat.get("label", "")), 14, theme["body"], theme["font"], align="center")
    return slide


def _build_quote_slide(prs, theme, quote_text, quote_author, background):
    bg = _bg_variant(theme, background or "accent")
    slide = _blank_slide(prs, theme, bg)
    fg = _contrast_text(bg)

    _add_textbox(slide, _in(0.8), _in(0.6), _in(2), _in(2), "\u201C", 90, fg, theme["font"], bold=True)
    _add_textbox(slide, _in(1.1), _in(2.3), prs.slide_width - _in(2.2), _in(2.2),
                 quote_text or "", 30, fg, theme["font"], italic=True, align="center")
    if quote_author:
        _add_textbox(slide, _in(1.1), _in(4.7), prs.slide_width - _in(2.2), _in(0.6),
                     f"— {quote_author}", 16, fg, theme["font"], align="center")
    return slide


def _build_process_slide(prs, theme, heading, steps, background, kicker):
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR_TYPE
    bg = _bg_variant(theme, background or "primary")
    slide = _blank_slide(prs, theme, bg)
    _add_accent_bar(slide, prs, theme["accent"], vertical=True)

    top = _in(0.5)
    if kicker:
        _add_kicker(slide, theme, kicker, _in(0.6), top, prs.slide_width - _in(1.2))
        top = top + _in(0.45)
    _add_textbox(slide, _in(0.6), top, prs.slide_width - _in(1.2), _in(0.8),
                 heading, 28, theme["title"], theme["font"], bold=True)

    steps = (steps or [])[:5]
    if not steps:
        return slide
    n = len(steps)
    margin = _in(0.9)
    circle_d = _in(0.7)
    row_y = _in(2.9)
    usable = prs.slide_width - margin * 2
    gap = usable / n

    centers_y = _i(row_y + circle_d / 2)
    line = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT,
                                       _i(margin + circle_d / 2), centers_y,
                                       _i(prs.slide_width - margin - circle_d / 2), centers_y)
    line.line.color.rgb = _rgb(theme["accent"])
    line.line.width = _pt(2)

    for i, step in enumerate(steps):
        slot_x = _i(margin + gap * i)
        cx = _i(slot_x + gap / 2 - circle_d / 2)
        circle = _add_shape(slide, MSO_SHAPE.OVAL, cx, row_y, circle_d, circle_d, fill_rgb=theme["accent"])
        tf = circle.text_frame
        tf.paragraphs[0].text = str(i + 1)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.size = _pt(20)
        tf.paragraphs[0].font.color.rgb = _rgb(_contrast_text(theme["accent"]))
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        _add_textbox(slide, slot_x, row_y + circle_d + _in(0.15), _i(gap), _in(0.5),
                     step.get("title", ""), 14, theme["title"], theme["font"], bold=True, align="center")
        _add_textbox(slide, slot_x, row_y + circle_d + _in(0.62), _i(gap), _in(1.3),
                     step.get("description", ""), 11, theme["body"], theme["font"], align="center")
    return slide


def _build_timeline_slide(prs, theme, heading, events, background, kicker):
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR_TYPE
    bg = _bg_variant(theme, background or "primary")
    slide = _blank_slide(prs, theme, bg)
    _add_accent_bar(slide, prs, theme["accent"], vertical=True)

    top = _in(0.5)
    if kicker:
        _add_kicker(slide, theme, kicker, _in(0.6), top, prs.slide_width - _in(1.2))
        top = top + _in(0.45)
    _add_textbox(slide, _in(0.6), top, prs.slide_width - _in(1.2), _in(0.8),
                 heading, 28, theme["title"], theme["font"], bold=True)

    events = (events or [])[:6]
    if not events:
        return slide
    n = len(events)
    margin = _in(0.9)
    dot_d = _in(0.22)
    mid_y = _in(3.7)
    usable = prs.slide_width - margin * 2
    gap = usable / n

    line = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT,
                                       _i(margin), _i(mid_y), _i(prs.slide_width - margin), _i(mid_y))
    line.line.color.rgb = _rgb(theme["accent"])
    line.line.width = _pt(2)

    for i, ev in enumerate(events):
        slot_x = _i(margin + gap * i)
        cx = _i(slot_x + gap / 2 - dot_d / 2)
        _add_shape(slide, MSO_SHAPE.OVAL, cx, _i(mid_y - dot_d / 2), dot_d, dot_d, fill_rgb=theme["accent"])
        _add_textbox(slide, slot_x, _i(mid_y - _in(0.9)), _i(gap), _in(0.6),
                     ev.get("date", ""), 14, theme["accent"], theme["font"], bold=True, align="center")
        _add_textbox(slide, slot_x, _i(mid_y + _in(0.3)), _i(gap), _in(1.0),
                     ev.get("label", ""), 12, theme["body"], theme["font"], align="center")
    return slide


def _place_cropped_image(slide, image_path, left, top, width, height):
    """Insert an image cropped to fill (left, top, width, height) without
    stretching/distorting it — the same 'cover fit' a real design tool does."""
    try:
        from PIL import Image
        with Image.open(image_path) as im:
            img_w, img_h = im.size
        pic = slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
        target_ratio = width / height
        img_ratio = img_w / img_h
        if img_ratio > target_ratio:
            crop = max(0.0, min(0.49, (1 - target_ratio / img_ratio) / 2))
            pic.crop_left = crop
            pic.crop_right = crop
        elif img_ratio < target_ratio:
            crop = max(0.0, min(0.49, (1 - img_ratio / target_ratio) / 2))
            pic.crop_top = crop
            pic.crop_bottom = crop
        return pic
    except Exception:
        try:
            return slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
        except Exception:
            return None


def _build_image_full_slide(prs, theme, heading, subheading, image_path, kicker):
    slide = _blank_slide(prs, theme, theme["bg"])
    if image_path:
        _place_cropped_image(slide, image_path, 0, 0, prs.slide_width, prs.slide_height)
    else:
        _set_background(slide, _bg_variant(theme, "muted"))

    from pptx.enum.shapes import MSO_SHAPE
    panel_h = _in(2.3)
    panel = _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, prs.slide_height - panel_h,
                        prs.slide_width, panel_h, fill_rgb=(0x00, 0x00, 0x00))
    _set_fill_alpha(panel, 55)

    text_top = prs.slide_height - panel_h + _in(0.4)
    if kicker:
        _add_textbox(slide, _in(0.7), text_top, prs.slide_width - _in(1.4), _in(0.4),
                     kicker.upper(), 12, theme["accent"], theme["font"], bold=True)
        text_top = text_top + _in(0.45)
    _add_textbox(slide, _in(0.7), text_top, prs.slide_width - _in(1.4), _in(0.9),
                 heading, 30, (0xFF, 0xFF, 0xFF), theme["font"], bold=True)
    if subheading:
        _add_textbox(slide, _in(0.7), text_top + _in(0.85), prs.slide_width - _in(1.4), _in(0.6),
                     subheading, 15, (0xE8, 0xE8, 0xE8), theme["font"])
    return slide


def _build_closing_slide(prs, theme, heading, subheading):
    bg = theme["accent"]
    slide = _blank_slide(prs, theme, bg)
    fg = _contrast_text(bg)
    _add_textbox(slide, _in(0.9), _in(2.9), prs.slide_width - _in(1.8), _in(1.4),
                 heading or "Thank You", 42, fg, theme["font"], bold=True, align="center")
    if subheading:
        _add_textbox(slide, _in(0.9), _in(4.1), prs.slide_width - _in(1.8), _in(0.8),
                     subheading, 16, fg, theme["font"], align="center")
    return slide


# ── orchestration ────────────────────────────────────────────────────────────

_LAYOUT_HANDLERS = {
    "bullets", "two_column", "table", "chart", "stat_highlight",
    "quote", "process", "timeline", "image_full", "section",
}


def _build_slide(prs, theme, entry: dict, idx: int, total: int, deck_title: str):
    layout = (entry.get("layout") or "bullets").strip().lower()
    if layout not in _LAYOUT_HANDLERS:
        layout = "bullets"
    heading = str(entry.get("heading", "")).strip() or "Untitled"
    background = (entry.get("background") or "").strip().lower() or None
    kicker = (entry.get("kicker") or "").strip()

    image_path = None
    wants_image = entry.get("image_path") or entry.get("image_query")

    if layout == "section":
        slide = _build_section_slide(prs, theme, heading, entry.get("subheading", ""), background, kicker)
    elif layout == "two_column":
        slide = _build_two_column_slide(
            prs, theme, heading,
            entry.get("left_heading", ""), entry.get("left_bullets") or [],
            entry.get("right_heading", ""), entry.get("right_bullets") or [],
            background, kicker,
        )
    elif layout == "table":
        slide = _build_table_slide(prs, theme, heading, entry.get("table") or {}, background, kicker)
    elif layout == "chart":
        slide = _build_chart_slide(prs, theme, heading, entry.get("chart") or {}, background, kicker)
    elif layout == "stat_highlight":
        slide = _build_stat_slide(prs, theme, heading, entry.get("stats") or [], background, kicker)
    elif layout == "quote":
        slide = _build_quote_slide(prs, theme, entry.get("quote_text", ""),
                                    entry.get("quote_author", ""), background)
    elif layout == "process":
        slide = _build_process_slide(prs, theme, heading, entry.get("process_steps") or [], background, kicker)
    elif layout == "timeline":
        slide = _build_timeline_slide(prs, theme, heading, entry.get("timeline_events") or [], background, kicker)
    elif layout == "image_full":
        if wants_image:
            image_path = _resolve_slide_image(entry)
        slide = _build_image_full_slide(prs, theme, heading, entry.get("subheading", ""), image_path, kicker)
    else:  # "bullets"
        if wants_image:
            image_path = _resolve_slide_image(entry)
        slide = _build_bullets_slide(prs, theme, heading, entry.get("bullets") or [],
                                      image_path, background, kicker)

    _add_footer(slide, prs, theme, deck_title, idx, total)

    notes = (entry.get("notes") or "").strip()
    if notes:
        try:
            slide.notes_slide.notes_text_frame.text = notes
        except Exception:
            pass

    return slide, layout, (wants_image is not None), (image_path is not None)


def run(parameters: dict, player=None, session_memory=None) -> str:
    title = (parameters.get("title") or "").strip()
    subtitle = (parameters.get("subtitle") or "").strip()
    author = (parameters.get("author") or "").strip()
    slides = parameters.get("slides") or []
    closing = parameters.get("closing") or {}
    output_path = (parameters.get("output_path") or "").strip()
    template_name = (parameters.get("template") or "modern_dark").strip()

    if not title:
        return "Sir, I need a title before I can build the presentation."
    if not slides:
        return "Sir, I need at least one content slide's worth of material to build the deck."

    if template_name not in THEMES:
        template_name = "modern_dark"
    theme = THEMES[template_name]

    try:
        from pptx import Presentation
    except ImportError:
        return ("Sir, the python-pptx library isn't installed. "
                "Run 'pip install python-pptx' and try again.")

    try:
        prs = Presentation()
        prs.slide_width = _in(13.333)   # widescreen 16:9
        prs.slide_height = _in(7.5)

        _build_title_slide(prs, theme, title, subtitle, author)

        total = len(slides) + (1 if closing else 0)
        images_used = 0
        images_failed = 0
        layouts_used = set()

        for i, entry in enumerate(slides, start=1):
            _, resolved_layout, wanted_image, got_image = _build_slide(prs, theme, entry, i, total, title)
            layouts_used.add(resolved_layout)
            if wanted_image:
                images_used += 1 if got_image else 0
                images_failed += 0 if got_image else 1

        if closing:
            _build_closing_slide(prs, theme, closing.get("heading", ""), closing.get("subheading", ""))

        save_path = _resolve_output_path(output_path, title)
        prs.save(str(save_path))

        result_text = (
            f"Presentation created: '{title}' — {len(slides)} content slide(s) across "
            f"{len(layouts_used)} layout type(s) ({', '.join(sorted(layouts_used))}), "
            f"'{template_name}' theme"
        )
        if images_used or images_failed:
            result_text += f", {images_used} image(s) placed"
            if images_failed:
                result_text += f" ({images_failed} couldn't be fetched, skipped)"
        result_text += f". Saved to {save_path}."

    except Exception as e:
        return f"Sir, create_powerpoint failed: {e}"

    if player:
        try:
            player.write_log(f"JARVIS: {result_text}")
        except Exception:
            pass

    return result_text
